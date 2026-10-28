from __future__ import annotations

import csv
import io
import re
from typing import Iterable

from app.phone import PhoneError, normalize_phone

_PHONEISH = re.compile(r"[\d+][\d\s\-()]{7,20}")
BULK_MAX_PHONES = 1000


def _cells_from_csv(data: bytes) -> Iterable[str]:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    for row in reader:
        for cell in row:
            yield str(cell).strip()


def _cells_from_xlsx(data: bytes) -> Iterable[str]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell is None:
                    continue
                yield str(cell).strip()
    finally:
        wb.close()


def parse_phone_file(*, filename: str, data: bytes) -> tuple[list[str], list[str]]:
    """Return (unique normalized phones, invalid raw values)."""
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        cells = _cells_from_xlsx(data)
    elif name.endswith(".csv") or name.endswith(".txt"):
        cells = _cells_from_csv(data)
    else:
        raise ValueError("file must be .csv or .xlsx")

    phones: list[str] = []
    seen: set[str] = set()
    invalid: list[str] = []

    for raw in cells:
        if not raw or raw.lower() in {"phone", "телефон", "tel", "mobile"}:
            continue
        if not _PHONEISH.search(raw) and not raw.isdigit():
            continue
        try:
            phone = normalize_phone(raw)
        except PhoneError:
            invalid.append(raw)
            continue
        if phone in seen:
            continue
        seen.add(phone)
        phones.append(phone)
        if len(phones) > BULK_MAX_PHONES:
            raise ValueError(f"max {BULK_MAX_PHONES} phones per file")

    if not phones:
        raise ValueError("no valid phones found in file")
    return phones, invalid
